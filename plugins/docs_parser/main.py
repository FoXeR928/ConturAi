from io import BytesIO
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from loguru import logger

from app.kernel import kernel
from app.kernel.contracts import BaseParser

# ==============================================================================
# БЛОК 1: КОНСТАНТЫ MIME-ТИПОВ (БЕЗ PDF)
# ==============================================================================
WORD_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "application/octet-stream",
}

EXCEL_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}

TEXT_MIMES = {
    "text/plain",
    "text/markdown",
    "text/x-markdown",
    "application/octet-stream",
}


# ==============================================================================
# БЛОК 2: ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ (ТАБЛИЦЫ И WORD)
# ==============================================================================

def _table_to_markdown(rows: list[list[str | None]]) -> str:
    if not rows:
        return ""
        
    rows = [r for r in rows if any(c and str(c).strip() for c in r)]
    if not rows:
        return ""
        
    width = max(len(r) for r in rows)
    norm = []
    for r in rows:
        r = [(str(c).strip() if c is not None else "") for c in r]
        r += [""] * (width - len(r))
        norm.append(r)
        
    header = "| " + " | ".join(norm[0]) + " |"
    sep = "| " + " | ".join(["---"] * width) + " |"
    body = ["| " + " | ".join(r) + " |" for r in norm[1:]]
    
    return "\n".join([header, sep] + body)


def _iter_block_items(doc):
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl

    parent = doc.element.body
    for child in parent.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


# ==============================================================================
# БЛОК 3: ПАРСИНГ WORD
# ==============================================================================

def read_word(data: bytes) -> str:
    try:
        doc = DocxDocument(BytesIO(data))
        parts: list[str] = []
        table_counter = 0

        for block in _iter_block_items(doc):
            if isinstance(block, Paragraph):
                t = block.text.strip()
                if t:
                    parts.append(t)
            elif isinstance(block, Table):
                table_counter += 1
                rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
                md = _table_to_markdown(rows)
                if md:
                    parts.append(f"--- Таблица {table_counter} ---\n{md}")

        result = "\n\n".join(parts)
        logger.debug(f"Word прочитан: {table_counter} таблиц")
        return result or "Документ не содержит текста"
    except Exception as err:
        logger.error(f"Word: {err}")
        return "Ошибка чтения файла"


# ==============================================================================
# БЛОК 4: ПАРСИНГ EXCEL
# ==============================================================================

def read_excel(data: bytes) -> str:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        
        for ws in wb.worksheets:
            parts.append(f"--- Лист: {ws.title} ---")
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                vals = [str(c).strip() if c is not None else "" for c in row]
                if any(vals):
                    rows.append(vals)
            md = _table_to_markdown(rows)
            if md:
                parts.append(md)
                
        result = "\n\n".join(parts)
        logger.debug(f"Excel прочитан: {len(wb.worksheets)} листов")
        return result or "Документ не содержит данных"
    except Exception as err:
        logger.error(f"Excel: {err}")
        return "Ошибка чтения файла"


# ==============================================================================
# БЛОК 5: ПАРСИНГ ТЕКСТА
# ==============================================================================

def read_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    logger.error("Не удалось декодировать текстовый файл")
    return "Ошибка чтения файла"


# ==============================================================================
# БЛОК 6: КЛАСС ПЛАГИНА
# ==============================================================================

@kernel.parser
class DocsParser(BaseParser):
    def can_handle(self, mime_type: str, filename: str) -> bool:
        ct = (mime_type or "").lower().split(";")[0].strip()
        name = (filename or "").lower()
        
        if ct in WORD_MIMES or name.endswith((".docx", ".doc")): 
            return True
        if ct in EXCEL_MIMES or name.endswith((".xlsx", ".xls")): 
            return True
        if ct in TEXT_MIMES or name.endswith((".txt", ".md", ".markdown", ".csv")): 
            return True
        return False

    def parse(self, mime_type: str, data: bytes, filename: str = "") -> str:
        ct = (mime_type or "").lower().split(";")[0].strip()
        name = (filename or "").lower()

        if ct in WORD_MIMES or name.endswith((".docx", ".doc")):
            return read_word(data)
        if ct in EXCEL_MIMES or name.endswith((".xlsx", ".xls")):
            return read_excel(data)
        if ct in TEXT_MIMES or name.endswith((".txt", ".md", ".markdown", ".csv")):
            return read_text(data)
        
        logger.warning(f"DocsParser: формат не поддерживается ({ct} / {name})")
        return "Ошибка: формат не поддерживается"