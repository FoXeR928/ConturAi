import pytesseract
from pdf2image import convert_from_bytes
from docx import Document
from openpyxl import load_workbook
from pymorphy3 import MorphAnalyzer
from re import sub
from langchain_text_splitters import RecursiveCharacterTextSplitter


def read_pdf(file):
    pages = convert_from_bytes(file.read(), 600)
    file_text = str()
    for page in pages:
        text = pytesseract.image_to_string(page, lang="rus")
        file_text += text + " "
    return file_text


def read_word(file):
    read = Document(file)
    file_text = str()
    if read.paragraphs:
        for paragraph in read.paragraphs:
            file_text += paragraph.text
    return file_text


def read_excel(file):
    dataTable = load_workbook(file)
    for sheet_name in dataTable.worksheets:
        sheet = dataTable[sheet_name.title]
        for row in sheet.iter_rows():
            for cell in row:
                file_text += f"{cell.value}. "
    return file_text


def normalize_text(text):
    morphanaliz = MorphAnalyzer()
    text = sub(r"[^\w\s]", "", text.lower())
    norm_text = " ".join(morphanaliz.normal_forms(word)[0] for word in text.split())
    return norm_text


def chunk_create(text):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=100,
        separators=["\n\n", "\n", "(?<=\.)!", " "],
        length_function=len,
    )
    chunks = splitter.split_text(text)
    return chunks
