from io import BytesIO

import pytesseract
from docx import Document as DocxDocument
from loguru import logger
from openpyxl import load_workbook
from pdf2image import convert_from_bytes

PDF_MIME = "application/pdf"
WORD_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def read_pdf(data: bytes) -> str:
    pages = convert_from_bytes(data, 600)
    return " ".join(pytesseract.image_to_string(p, lang="rus") for p in pages)


def read_word(data: bytes) -> str:
    doc = DocxDocument(BytesIO(data))
    return " ".join(p.text for p in doc.paragraphs)


def read_excel(data: bytes) -> str:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts.append(" ".join(str(c) for c in row if c is not None))
    return " ".join(parts)


def parse_file(content_type: str, data: bytes) -> str | None:
    try:
        if content_type == PDF_MIME:
            return read_pdf(data)
        if content_type == WORD_MIME:
            return read_word(data)
        if content_type == EXCEL_MIME:
            return read_excel(data)
        return None
    except Exception as err:
        logger.error(f"Не удалось прочитать документ: {err}")
        return None